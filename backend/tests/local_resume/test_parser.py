import asyncio
import io
import unittest
from unittest.mock import patch
from resume_parser import parse_resume_text, extract_text_from_file

SAMPLE = '''Jane Example, RN
jane@example.com | (555) 321-9876
Austin, TX 78701
Professional Summary
Registered nurse with patient care and wound care experience.
Work Experience
Registered Nurse | Example Hospital
January 2021 - Present
- Provide wound care and record vital signs.
Education
Example University
Bachelor of Science in Nursing 2020
Certifications
BLS
RN License
Skills
Epic; Scheduling
'''

class ParserTests(unittest.TestCase):
    def test_extract_without_api_or_network(self):
        async def run_offline():
            with patch.dict('os.environ', {}, clear=True), patch('socket.socket', side_effect=AssertionError('network forbidden')):
                return await parse_resume_text(SAMPLE)
        data = asyncio.run(run_offline())
        self.assertEqual(data['full_name'], 'Jane Example')
        self.assertEqual(data['email'], 'jane@example.com')
        self.assertEqual(data['address']['state'], 'TX')
        self.assertEqual(data['employment_groups'][0]['start_date'], '2021-01')
        self.assertTrue(data['employment_groups'][0]['currently_working'])
        self.assertEqual(data['education'][0]['institution'], 'Example University')
        self.assertEqual(data['certifications'][0]['name'], 'BLS')
        self.assertEqual(data['_parse_metadata']['method'], 'local_rules')

    def test_no_license_invented_from_job_title(self):
        data = asyncio.run(parse_resume_text(SAMPLE.replace('RN License', '')))
        self.assertEqual(data['licenses'], [])

    def test_text_pdf_and_docx_table(self):
        import fitz
        import docx
        pdf = fitz.open(); page = pdf.new_page(); page.insert_text((40, 40), SAMPLE, fontsize=9)
        pdf_bytes = pdf.tobytes(); pdf.close()
        document = docx.Document(); document.add_table(rows=1, cols=1).cell(0, 0).text = SAMPLE
        stream = io.BytesIO(); document.save(stream)
        for filename, contents in [('resume.pdf', pdf_bytes), ('resume.docx', stream.getvalue()), ('resume.txt', SAMPLE.encode())]:
            text = asyncio.run(extract_text_from_file(contents, filename))
            self.assertIn('jane@example.com', text)
            self.assertNotIn('error', asyncio.run(parse_resume_text(text)))

    def test_empty_scan_and_corrupt_file_are_actionable(self):
        import fitz
        doc = fitz.open(); doc.new_page(); contents = doc.tobytes(); doc.close()
        self.assertIn('Scanned PDFs', asyncio.run(extract_text_from_file(contents, 'scan.pdf')))
        self.assertIn('could not be read', asyncio.run(extract_text_from_file(b'bad', 'bad.docx')))
        self.assertIn('Older .doc', asyncio.run(extract_text_from_file(b'old', 'old.doc')))

if __name__ == '__main__': unittest.main()
